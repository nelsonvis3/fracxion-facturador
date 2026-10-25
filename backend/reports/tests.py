import csv
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from unittest.mock import patch

from django.core.cache import cache
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from auditlog.models import AuditEvent
from billing.services import record_payment, reverse_payment
from billing.test_regressions import BillingFixtures
from tenancy.models import Tenant

from .services import ReportTooLarge, dataset, render_report, safe_text


class ReportTests(BillingFixtures, TestCase):
    def setUp(self):
        super().setUp()
        cache.clear()
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.url = f"/api/tenants/{self.tenant.pk}/reports/export/"
        self.filters = {"date_from": "2026-10-01", "date_to": "2026-10-31", "currency": "ARS"}
        with patch(
            "billing.services.timezone.now", return_value=datetime(2026, 10, 14, 15, tzinfo=UTC)
        ):
            self.invoice = self.issued()

    def payment(self):
        return record_payment(
            tenant=self.tenant,
            invoice_id=self.invoice.pk,
            actor=self.user,
            amount="40.00",
            method="cash",
            received_at=datetime(2026, 10, 14, 16, tzinfo=UTC),
        )

    def test_all_report_types_and_formats_contain_real_company_metadata(self):
        self.payment()
        for kind in ["sales", "payments", "receivables", "products", "summary", "customers"]:
            data = dataset(self.tenant, self.filters | {"kind": kind})
            self.assertTrue(data["rows"])
            for file_format in ["csv", "xlsx", "pdf"]:
                result = render_report(data, file_format)
                self.assertGreater(len(result), 200)
                if file_format == "pdf":
                    self.assertTrue(result.startswith(b"%PDF-"))
                elif file_format == "xlsx":
                    sheet = load_workbook(BytesIO(result), data_only=False).active
                    self.assertIn(self.tenant.name, [cell.value for row in sheet for cell in row])
                    self.assertTrue(all(cell.data_type != "f" for row in sheet for cell in row))
                else:
                    parsed = list(csv.reader(StringIO(result.decode("utf-8-sig"))))
                    self.assertIn(["Empresa", self.tenant.name], parsed)
                    self.assertIn(["Moneda", "ARS"], parsed)

    def test_formula_payloads_are_literal_and_money_stays_exact(self):
        for payload in [
            '=WEBSERVICE("https://example.invalid")',
            " \t+SUM(1,1)",
            "-1+1",
            "@SUM(1)",
            "＝SUM(1)",
            "\n=1+1",
        ]:
            self.assertTrue(safe_text(payload).startswith("'"))
        self.assertEqual(safe_text(Decimal("-12345678901234.12")), "-12345678901234.12")
        self.assertEqual(safe_text("safe\x00text"), "safetext")
        self.customer.name = "=1+1"
        self.customer.save()
        # Historical descriptions remain the immutable source for issued sales.
        self.assertEqual(dataset(self.tenant, self.filters)["rows"][0][1], "Customer")
        sample = dataset(self.tenant, self.filters)
        sample["rows"][0][1] = "=1+1"
        sheet = load_workbook(BytesIO(render_report(sample, "xlsx")), data_only=False).active
        self.assertTrue(any(cell.value == "'=1+1" for row in sheet for cell in row))
        self.assertFalse(any(cell.data_type == "f" for row in sheet for cell in row))

    def test_payment_reversal_keeps_original_record_and_zero_effective_amount(self):
        payment = self.payment()
        reverse_payment(
            tenant=self.tenant, payment_id=payment.pk, actor=self.user, reason="Correction"
        )
        result = dataset(self.tenant, self.filters | {"kind": "payments"})
        self.assertEqual(result["rows"][0][-3:], [Decimal("40.00"), Decimal("0.00"), "Revertido"])
        ledger = dataset(self.tenant, self.filters | {"kind": "customers"})
        self.assertEqual(sum(row[4] for row in ledger["rows"]), Decimal("100.00"))

    def test_filters_empty_period_and_limits(self):
        self.assertEqual(dataset(self.tenant, self.filters | {"currency": "USD"})["rows"], [])
        self.assertEqual(
            dataset(
                self.tenant, self.filters | {"customer": "00000000-0000-4000-8000-000000000001"}
            )["rows"],
            [],
        )
        with self.assertRaises(ReportTooLarge):
            dataset(self.tenant, self.filters, max_rows=0)
        self.assertEqual(
            self.client.get(self.url, self.filters | {"file_format": "exe"}).status_code, 400
        )

    def test_http_permissions_isolation_and_export_audit(self):
        for role in ["owner", "admin", "accountant", "sales", "viewer"]:
            self.membership.role = role
            self.membership.save()
            response = self.client.get(self.url, self.filters)
            self.assertEqual(
                response.status_code, 200 if role in {"owner", "admin", "accountant"} else 403
            )
            if response.status_code == 200:
                self.assertIn("no-store", response["Cache-Control"])
                self.assertIn("attachment", response["Content-Disposition"])
        self.assertEqual(
            AuditEvent.objects.filter(tenant=self.tenant, entity_type="reports.export").count(), 3
        )
        foreign = Tenant.objects.create(name="Other", slug="other-reports")
        self.assertEqual(
            self.client.get(f"/api/tenants/{foreign.pk}/reports/export/").status_code, 403
        )
        self.membership.role = "owner"
        self.membership.is_active = False
        self.membership.save()
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.membership.is_active = True
        self.membership.save()
        Tenant.objects.filter(pk=self.tenant.pk).update(is_active=False)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get(self.url).status_code, 401)

    def test_rate_limit_cannot_be_bypassed_with_spoofed_headers(self):
        for _ in range(10):
            self.assertEqual(self.client.get(self.url, self.filters).status_code, 200)
        self.assertEqual(
            self.client.get(self.url, self.filters, HTTP_X_FORWARDED_FOR="203.0.113.4").status_code,
            429,
        )

    def test_empty_and_long_text_pdf_paginate_without_fetching_urls(self):
        data = dataset(self.tenant, self.filters)
        data["rows"][0][1] = "<img src='http://127.0.0.1:9/private'/>" + "detalle " * 2000
        result = render_report(data, "pdf")
        self.assertTrue(result.startswith(b"%PDF-"))
        data["rows"] = []
        self.assertTrue(render_report(data, "pdf").startswith(b"%PDF-"))

    def test_expense_reports_preserve_scope_currency_and_reversed_payments(self):
        from operations.expenses import (
            pay_expense,
            post_expense,
            reverse_expense_payment,
            save_expense,
        )
        from operations.models import Supplier

        supplier = Supplier.objects.create(tenant=self.tenant, name="Vendor reporting")
        expense = save_expense(
            tenant_id=self.tenant.pk,
            actor=self.user,
            data={
                "description": "Actual synthetic expense",
                "incurred_on": "2026-10-14",
                "currency": "ARS",
                "net_amount": "100.00",
                "tax_rate": "21.00",
                "supplier": str(supplier.pk),
            },
        )
        post_expense(tenant_id=self.tenant.pk, actor=self.user, expense_id=expense.pk)
        payment = pay_expense(
            tenant_id=self.tenant.pk,
            actor=self.user,
            expense_id=expense.pk,
            data={"amount": "40.00", "method": "cash"},
        )
        data = dataset(
            self.tenant, self.filters | {"kind": "expenses", "supplier": str(supplier.pk)}
        )
        self.assertEqual(
            data["rows"][0][-3:], [Decimal("121.00"), Decimal("40.00"), Decimal("81.00")]
        )
        for file_format in ["csv", "xlsx", "pdf"]:
            self.assertGreater(len(render_report(data, file_format)), 200)
        reverse_expense_payment(
            tenant_id=self.tenant.pk,
            actor=self.user,
            payment_id=payment.pk,
            data={"reason": "correction"},
        )
        self.assertEqual(
            dataset(self.tenant, self.filters | {"kind": "payables"})["rows"][0][-1],
            Decimal("121.00"),
        )
        self.assertEqual(
            dataset(self.tenant, self.filters | {"kind": "expenses", "currency": "USD"})["rows"], []
        )
        from rest_framework.exceptions import ValidationError

        with self.assertRaises(ValidationError):
            dataset(
                self.tenant, self.filters | {"kind": "expenses", "customer": str(self.customer.pk)}
            )
