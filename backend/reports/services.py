import csv
import re
import unicodedata
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from xml.sax.saxutils import escape

import reportlab
from django.db.models import Sum
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
from rest_framework import serializers

from analytics.services import (
    AnalyticsFilters,
    boundaries,
    period_metrics,
    report_queries,
    reporting_context,
)
from analytics.services import amount as exact_amount
from analytics.snapshot import reporting_snapshot
from billing.models import InvoiceLine, Payment, PaymentReversal
from tenancy.serializers import StrictInputMixin

REPORT_TITLES = {
    "sales": "Resumen de ventas internas",
    "payments": "Informe de cobros",
    "receivables": "Cuentas por cobrar",
    "customers": "Movimientos por cliente",
    "products": "Facturación por producto",
    "summary": "Resumen financiero interno",
    "expenses": "Gastos registrados",
    "payables": "Cuentas por pagar",
}
MIME_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


class ReportFilters(StrictInputMixin, AnalyticsFilters):
    kind = serializers.ChoiceField(choices=list(REPORT_TITLES), default="sales")
    file_format = serializers.ChoiceField(choices=list(MIME_TYPES), default="csv")
    supplier = serializers.UUIDField(required=False)

    def validate(self, attrs):
        expenses = attrs.get("kind", "sales") in {"expenses", "payables"}
        if expenses and attrs.get("customer"):
            raise serializers.ValidationError(
                {"customer": "El informe de gastos se filtra por proveedor, no por cliente."}
            )
        if not expenses and attrs.get("supplier"):
            raise serializers.ValidationError(
                {"supplier": "Este informe no admite filtro de proveedor."}
            )
        return attrs


class ReportTooLarge(Exception):
    """The caller must enqueue a durable background export."""


def money_cell(value):
    return Decimal(exact_amount(value))


def payment_records(tenant, values, zone):
    start, end = boundaries(values["date_from"], values["date_to"], zone)
    query = Payment.objects.filter(
        tenant=tenant,
        invoice__tenant=tenant,
        invoice__currency=values["currency"],
        invoice__status="issued",
        invoice__document_kind="invoice",
        received_at__gte=start,
        received_at__lt=end,
    ).select_related("invoice", "reversal")
    return (
        query.filter(invoice__customer_id=values["customer"]) if values.get("customer") else query
    )


def dataset(tenant, params, *, max_rows=1000):
    validator = ReportFilters(data=params)
    validator.is_valid(raise_exception=True)
    kind = validator.validated_data["kind"]
    with reporting_snapshot():
        values, zone, today = reporting_context(tenant, params)
        invoices, payments = report_queries(tenant, values, zone)

        def limited(query):
            rows = list(query[: max_rows + 1])
            if len(rows) > max_rows:
                raise ReportTooLarge()
            return rows

        if kind in {"expenses", "payables"}:
            from operations.expenses import expense_query

            query = expense_query(tenant).filter(
                status="posted",
                currency=values["currency"],
                incurred_on__gte=values["date_from"],
                incurred_on__lte=values["date_to"],
            )
            if validator.validated_data.get("supplier"):
                query = query.filter(supplier_id=validator.validated_data["supplier"])
            if kind == "payables":
                query = query.filter(balance__gt=0)
            headers = [
                "ID gasto",
                "Proveedor histórico",
                "Categoría histórica",
                "Descripción",
                "Referencia",
                "Fecha gasto",
                "Vencimiento",
                "Neto",
                "Impuesto",
                "Total",
                "Pagado efectivo",
                "Saldo actual",
            ]
            rows = [
                [
                    str(row.pk),
                    row.supplier_name_snapshot,
                    row.category_name_snapshot,
                    row.description,
                    row.reference,
                    str(row.incurred_on),
                    str(row.due_date or ""),
                    money_cell(row.net_amount),
                    money_cell(row.tax_amount),
                    money_cell(row.total),
                    money_cell(row.paid_amount),
                    money_cell(row.balance),
                ]
                for row in limited(query.order_by("incurred_on", "pk"))
            ]
        elif kind == "sales":
            headers = [
                "Número interno",
                "Cliente histórico",
                "Fecha emisión",
                "Vencimiento",
                "Subtotal",
                "Impuestos",
                "Total",
                "Cobrado efectivo",
                "Saldo actual",
            ]
            rows = [
                [
                    row.internal_number,
                    row.customer_name_snapshot,
                    row.issued_at.astimezone(zone).isoformat(),
                    str(row.due_date or ""),
                    money_cell(row.subtotal),
                    money_cell(row.tax_total),
                    money_cell(row.total),
                    money_cell(row.paid_amount),
                    money_cell(row.current_balance),
                ]
                for row in limited(invoices.order_by("issued_at", "pk"))
            ]
        elif kind == "receivables":
            headers = [
                "Número interno",
                "Cliente histórico",
                "Fecha emisión",
                "Vencimiento",
                "Días de atraso",
                "Total",
                "Saldo actual",
            ]
            rows = [
                [
                    row.internal_number,
                    row.customer_name_snapshot,
                    row.issued_at.astimezone(zone).isoformat(),
                    str(row.due_date or ""),
                    max((today - row.due_date).days, 0) if row.due_date else "Sin fecha",
                    money_cell(row.total),
                    money_cell(row.current_balance),
                ]
                for row in limited(
                    invoices.filter(current_balance__gt=0).order_by("due_date", "pk")
                )
            ]
        elif kind == "payments":
            headers = [
                "Pago",
                "Factura interna",
                "Cliente histórico",
                "Recepción",
                "Medio",
                "Referencia",
                "Importe original",
                "Importe efectivo",
                "Estado",
            ]
            rows = []
            for row in limited(payment_records(tenant, values, zone).order_by("received_at", "pk")):
                reversed_payment = hasattr(row, "reversal")
                rows.append(
                    [
                        str(row.pk),
                        row.invoice.internal_number,
                        row.invoice.customer_name_snapshot,
                        row.received_at.astimezone(zone).isoformat(),
                        row.get_method_display(),
                        row.external_reference or "",
                        money_cell(row.amount),
                        money_cell(0 if reversed_payment else row.amount),
                        "Revertido" if reversed_payment else "Vigente",
                    ]
                )
        elif kind == "products":
            headers = [
                "Producto",
                "Descripción histórica",
                "SKU histórico",
                "Cantidad",
                "Subtotal",
                "Impuestos",
                "Total",
            ]
            query = (
                InvoiceLine.objects.filter(invoice__in=invoices.values("pk"))
                .values("product_id", "description_snapshot", "sku_snapshot")
                .annotate(
                    quantity=Sum("quantity"),
                    subtotal=Sum("subtotal"),
                    tax=Sum("tax_amount"),
                    total=Sum("total"),
                )
                .order_by("description_snapshot", "product_id")
            )
            rows = [
                [
                    str(r["product_id"] or ""),
                    r["description_snapshot"],
                    r["sku_snapshot"],
                    str(r["quantity"]),
                    money_cell(r["subtotal"]),
                    money_cell(r["tax"]),
                    money_cell(r["total"]),
                ]
                for r in limited(query)
            ]
        elif kind == "summary":
            metrics = period_metrics(invoices, payments, today)
            headers = ["Indicador", "Valor", "Criterio"]
            rows = [
                ["Facturado", metrics["invoiced"], "Por fecha de emisión"],
                [
                    "Cobrado efectivo",
                    metrics["collected"],
                    "Por fecha de recepción; excluye reversiones",
                ],
                ["Saldo actual", metrics["outstanding"], "Facturas emitidas en el período"],
                [
                    "Saldo vencido actual",
                    metrics["overdue"],
                    "Con fecha válida; facturas del período",
                ],
                ["Facturas emitidas", metrics["issued_count"], "Cantidad"],
                ["Facturas pendientes", metrics["unpaid_count"], "Saldo actual positivo"],
            ]
        else:
            headers = [
                "Cliente histórico",
                "Fecha registrada",
                "Movimiento",
                "Documento",
                "Importe firmado",
                "Referencia",
            ]
            events = []
            for row in limited(invoices.order_by("issued_at", "pk")):
                events.append(
                    (
                        row.issued_at,
                        str(row.pk),
                        [
                            row.customer_name_snapshot,
                            row.issued_at.astimezone(zone).isoformat(),
                            "Emisión",
                            str(row.internal_number),
                            money_cell(row.total),
                            "",
                        ],
                    )
                )
            for row in limited(payment_records(tenant, values, zone).order_by("received_at", "pk")):
                events.append(
                    (
                        row.received_at,
                        str(row.pk),
                        [
                            row.invoice.customer_name_snapshot,
                            row.received_at.astimezone(zone).isoformat(),
                            "Cobro",
                            str(row.invoice.internal_number),
                            money_cell(-row.amount),
                            row.external_reference or "",
                        ],
                    )
                )
            start, end = boundaries(values["date_from"], values["date_to"], zone)
            reversals = PaymentReversal.objects.filter(
                tenant=tenant,
                payment__tenant=tenant,
                payment__invoice__tenant=tenant,
                payment__invoice__currency=values["currency"],
                payment__invoice__status="issued",
                payment__invoice__document_kind="invoice",
                occurred_at__gte=start,
                occurred_at__lt=end,
            ).select_related("payment__invoice")
            if values.get("customer"):
                reversals = reversals.filter(payment__invoice__customer_id=values["customer"])
            for row in limited(reversals.order_by("occurred_at", "pk")):
                events.append(
                    (
                        row.occurred_at,
                        str(row.pk),
                        [
                            row.payment.invoice.customer_name_snapshot,
                            row.occurred_at.astimezone(zone).isoformat(),
                            "Reversión",
                            str(row.payment.invoice.internal_number),
                            money_cell(row.payment.amount),
                            row.reason,
                        ],
                    )
                )
            if len(events) > max_rows:
                raise ReportTooLarge()
            rows = [row for _, _, row in sorted(events, key=lambda item: (item[0], item[1]))]
        metadata = [
            ["Informe", REPORT_TITLES[kind]],
            ["Empresa", tenant.name],
            ["ID empresa", str(tenant.pk)],
            ["Desde", str(values["date_from"])],
            ["Hasta", str(values["date_to"])],
            ["Moneda", values["currency"]],
            ["Zona horaria", str(zone)],
            ["Generado", timezone.now().astimezone(zone).isoformat()],
            [
                "Criterio",
                "Gastos registrados por fecha del gasto; pagos y saldos actuales. No constituye cierre contable ni valida comprobantes externos."
                if kind in {"expenses", "payables"}
                else "Datos actuales; saldos de facturas emitidas en el período. Comprobantes internos sin autorización fiscal de ARCA.",
            ],
        ]
        if kind == "customers":
            metadata.append(
                [
                    "Movimientos",
                    "Fechas registradas; no incluye saldo inicial ni constituye cierre contable.",
                ]
            )
        return {
            "kind": kind,
            "metadata": metadata,
            "headers": headers,
            "rows": rows,
            "currency": values["currency"],
            "date_from": str(values["date_from"]),
            "date_to": str(values["date_to"]),
        }


def safe_text(value):
    if isinstance(value, Decimal):
        return format(value, ".2f")
    text = str(value if value is not None else "")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    if unicodedata.normalize("NFKC", text).lstrip().startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text


def render_report(data, file_format):
    if file_format not in MIME_TYPES:
        raise serializers.ValidationError({"file_format": "Formato no admitido."})
    if file_format == "csv":
        output = StringIO(newline="")
        writer = csv.writer(output)
        for row in [*data["metadata"], [], data["headers"], *data["rows"]]:
            writer.writerow([safe_text(value) for value in row])
        return output.getvalue().encode("utf-8-sig")
    if file_format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Informe"
        sheet.append(["Importes decimales exactos como texto; sin conversión de moneda."])
        for row in data["metadata"]:
            sheet.append([safe_text(value) for value in row])
        sheet.append([])
        sheet.append(data["headers"])
        header_row = sheet.max_row
        for row in data["rows"]:
            sheet.append([safe_text(value) for value in row])
        for row in sheet.iter_rows():
            for cell in row:
                cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4338CA")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 26
        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = (
            f"A{header_row}:{sheet.cell(sheet.max_row, len(data['headers'])).coordinate}"
        )
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    # Only generated text and a bundled font; no external images/URLs or user markup.
    font_path = Path(reportlab.__file__).parent / "fonts/Vera.ttf"
    if "FracxionVera" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("FracxionVera", str(font_path)))
    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=28,
        rightMargin=28,
        topMargin=28,
        bottomMargin=32,
        title=REPORT_TITLES[data["kind"]],
        author="Fracxion",
    )
    body = ParagraphStyle(
        "ReportBody",
        fontName="FracxionVera",
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor("#1e293b"),
        splitLongWords=True,
    )
    heading = ParagraphStyle("ReportHeading", parent=body, fontSize=15, leading=20, spaceAfter=12)
    header = ParagraphStyle("ReportTableHeading", parent=body, textColor=colors.white)
    story = [Paragraph(escape(REPORT_TITLES[data["kind"]]), heading)]
    for key, value in data["metadata"]:
        story.append(Paragraph(escape(f"{key}: {value}"), body))
    story.append(Spacer(1, 12))
    table_data = [[Paragraph(escape(str(value)), header) for value in data["headers"]]]
    table_data.extend(
        [[Paragraph(escape(str(value)), body) for value in row] for row in data["rows"]]
    )
    table = LongTable(
        table_data,
        repeatRows=1,
        splitInRow=1,
        colWidths=[doc.width / len(data["headers"])] * len(data["headers"]),
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4338ca")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor("#cbd5e1")),
            ]
        )
    )
    story.append(table)

    def page_number(canvas, document):
        canvas.saveState()
        canvas.setFont("FracxionVera", 8)
        canvas.drawRightString(document.pagesize[0] - 28, 16, f"Fracxion · Página {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=page_number, onLaterPages=page_number)
    return output.getvalue()
